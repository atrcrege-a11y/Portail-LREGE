# -*- coding: utf-8 -*-
"""Lot 3 — générateur du classeur depuis publication.json.

Le critère de fin du §4 est une recette, pas une assertion de plus : le
classeur régénéré depuis la lecture du classeur réel doit lui être identique
cellule par cellule, aux écarts documentés près. Ces écarts sont énumérés un
par un — c'est leur liste, et non leur nombre, qui fait la recette.

Depuis le Lot 4, la comparaison elle-même et la liste des écarts attendus
vivent dans `recette_classeur.py` : ce test les importe de là plutôt que d'en
tenir une copie, qui dériverait. La normalisation (décision 7 du §5 bis —
espaces de bord, cellules saisies en vraie date Excel comme `Épée!A13`) y est
faite aussi.
"""
import datetime
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import export_classeur as ec
import import_classeur as ic
import publication as pub
from recette_classeur import (ECARTS_ATTENDUS, ECARTS_ATTENDUS_CORRIGES,
                              comparer)
from texte_lrege import txt

FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures')
CLASSEUR = os.environ.get('LREGE_CLASSEUR') or os.path.join(
    FIXTURES, 'Calendrier_LREGE_2026-2027.xlsx')

# ─────────────────────────────────────────────────────────────────────────────
# Classeurs régénérés
# ─────────────────────────────────────────────────────────────────────────────

@pytest.fixture(scope='module')
def regenere(tmp_path_factory):
    """Classeur régénéré depuis la lecture **sans** les corrections du §7.

    C'est cette lecture-là que la recette compare à l'original : les
    corrections en attente sont, par construction, des écarts voulus.
    """
    lignes, _ = ic.lire_classeur(CLASSEUR, appliquer_corrections=False)
    for l in lignes:
        l.pop('_source', None)
    chemin = str(tmp_path_factory.mktemp('lot3') / 'regenere.xlsx')
    rapport = ec.ecrire_classeur(lignes, chemin)
    return chemin, lignes, rapport


@pytest.fixture(scope='module')
def regenere_corrige(tmp_path_factory):
    lignes, _ = ic.lire_classeur(CLASSEUR, appliquer_corrections=True)
    for l in lignes:
        l.pop('_source', None)
    chemin = str(tmp_path_factory.mktemp('lot3c') / 'regenere.xlsx')
    ec.ecrire_classeur(lignes, chemin)
    return chemin, lignes


# ─────────────────────────────────────────────────────────────────────────────
# Critère de fin du Lot 3 (§4 du plan)
# ─────────────────────────────────────────────────────────────────────────────

class TestCritereDeFin:

    def test_ecarts_exactement_ceux_attendus(self, regenere):
        """Le classeur régénéré est identique à l'original, écarts documentés mis à part."""
        assert comparer(CLASSEUR, regenere[0]) == ECARTS_ATTENDUS

    def test_douze_onglets_sans_lisez_moi(self, regenere):
        wb = openpyxl.load_workbook(regenere[0])
        assert wb.sheetnames == [ic.ONGLET_CHRONO] + ic.ONGLETS_ARME \
            + [n for n, _ in ic.ONGLETS_TOURNOI]
        assert len(wb.sheetnames) == 12
        assert ic.ONGLET_IGNORE not in wb.sheetnames

    def test_avec_corrections_les_ecarts_en_plus_sont_ceux_du_paragraphe_7(
            self, regenere_corrige):
        """Avec les corrections, l'écart est celui-ci et rien d'autre."""
        obtenus = comparer(CLASSEUR, regenere_corrige[0])
        assert obtenus == ECARTS_ATTENDUS_CORRIGES
        en_plus = [e for e in obtenus if e not in ECARTS_ATTENDUS]
        assert len(en_plus) == 8, en_plus


# ─────────────────────────────────────────────────────────────────────────────
# Structure du classeur produit
# ─────────────────────────────────────────────────────────────────────────────

class TestStructure:

    def test_entetes_et_colonnes_supprimees(self, regenere):
        wb = openpyxl.load_workbook(regenere[0])
        chrono = wb[ic.ONGLET_CHRONO]
        assert [chrono.cell(2, c).value for c in range(1, 9)] == ec.ENTETES_CHRONO
        assert all(chrono.cell(r, 9).value is None
                   for r in range(1, chrono.max_row + 1)), 'colonne I non vidée'
        for nom in ic.ONGLETS_ARME:
            assert [wb[nom].cell(2, c).value for c in range(1, 6)] == ec.ENTETES_ARME
        for nom, _ in ic.ONGLETS_TOURNOI:
            ws = wb[nom]
            assert [ws.cell(2, c).value for c in range(1, 13)] == ec.ENTETES_TOURNOI
            assert all(ws.cell(r, 13).value is None
                       for r in range(1, ws.max_row + 1)), 'colonne M non vidée'

    def test_titres_et_saison_deduite(self, regenere):
        wb = openpyxl.load_workbook(regenere[0])
        assert regenere[2]['saison'] == '2026-2027'
        assert wb[ic.ONGLET_CHRONO].cell(1, 1).value == ec.TITRE_CHRONO % '2026-2027'
        assert wb['Sabre laser'].cell(1, 1).value == 'SABRE LASER — saison 2026-2027'
        assert wb['Tournois Champagne-Ardenne'].cell(1, 1).value \
            == 'TOURNOIS DE CLUBS — CHAMPAGNE-ARDENNE — saison 2026-2027'

    def test_aucune_cellule_datee(self, regenere):
        """Dates en texte : une cellule datée se relit autrement d'un poste à l'autre."""
        wb = openpyxl.load_workbook(regenere[0])
        datees = [(ws.title, c.coordinate) for ws in wb.worksheets
                  for row in ws.iter_rows() for c in row
                  if isinstance(c.value, (datetime.date, datetime.datetime))]
        assert datees == []

    def test_les_donnees_commencent_ligne_3(self, regenere):
        wb = openpyxl.load_workbook(regenere[0])
        for nom in wb.sheetnames:
            ws = wb[nom]
            assert ws.cell(1, 1).value, '%s : titre manquant' % nom
            assert ws.cell(2, 1).value == 'Date'


class TestOngletsDerives:

    def test_onglet_par_arme_derive_du_champ_armes(self, regenere):
        """Zone et Master paraissent sur l'onglet de chacune de leurs armes."""
        wb = openpyxl.load_workbook(regenere[0])
        types = {nom: [wb[nom].cell(r, 3).value
                       for r in range(3, wb[nom].max_row + 1)
                       if wb[nom].cell(r, 3).value]
                 for nom in ic.ARMES}
        assert types['Épée'][-2:] == ['Zone', 'Master']
        assert types['Fleuret'][-2:] == ['Zone', 'Master']
        assert types['Sabre'][-1] == 'Master' and 'Zone' not in types['Sabre']
        assert 'Master' not in types['Sabre laser'] and 'Master' not in types['Para']

    def test_master_repris_au_23_mai_partout(self, regenere):
        """Décision 1 du §5 : les onglets par arme font foi, Chronologique compris."""
        wb = openpyxl.load_workbook(regenere[0])
        chrono = wb[ic.ONGLET_CHRONO]
        derniere = chrono.max_row
        assert [chrono.cell(derniere, c).value for c in range(1, 5)] \
            == ['23 mai 2027', 'dim', 'Transverse', 'Master']
        for nom in ('Épée', 'Fleuret', 'Sabre'):
            ws = wb[nom]
            assert ws.cell(ws.max_row, 1).value == '23 mai 2027'
            assert ws.cell(ws.max_row, 5).value == 'Wassy'

    def test_ligne_vide_avant_chaque_ligne_transverse(self, regenere):
        ws = openpyxl.load_workbook(regenere[0])['Épée']
        assert [ws.cell(r, 1).value for r in (10, 12)] == [None, None]
        assert ws.cell(11, 3).value == 'Zone'
        assert ws.cell(13, 1).value == '23 mai 2027'

    def test_onglet_transverse_exactement_les_coupes_de_lorraine(self, regenere):
        ws = openpyxl.load_workbook(regenere[0])['Transverse']
        types = [ws.cell(r, 3).value for r in range(3, ws.max_row + 1)]
        assert types == ['CDL 1', 'CDL 2', 'CDL 3', 'CHDL']
        assert 'Master' not in types and 'CID' not in types

    def test_lignes_de_tournoi_du_chronologique_derivees(self, regenere):
        """La restitution passe par `rendu_chronologique_tournoi()`, jamais par les champs."""
        chemin, lignes, _ = regenere
        ws = openpyxl.load_workbook(chemin)[ic.ONGLET_CHRONO]
        lues = [tuple(txt(ws.cell(r, c).value) for c in range(1, 9))
                for r in range(3, ws.max_row + 1)
                if txt(ws.cell(r, 3).value).startswith('Tournoi ')]
        attendu = [tuple(txt(v) for v in ic.rendu_chronologique_tournoi(l))
                   for l in ec._ordonnees(lignes) if l['est_tournoi']]
        assert lues == attendu and len(lues) == 9

    def test_ligne_de_total_des_onglets_tournois(self, regenere):
        wb = openpyxl.load_workbook(regenere[0])
        alsace = wb['Tournois Alsace']
        assert alsace.cell(9, 1).value is None            # séparation
        assert alsace.cell(10, 1).value == '6 tournoi(s) déclaré(s).'
        lorraine = wb['Tournois Lorraine']
        assert lorraine.cell(7, 1).value == '3 tournoi(s) déclaré(s).'

    def test_territoire_sans_tournoi(self, regenere):
        ws = openpyxl.load_workbook(regenere[0])['Tournois Champagne-Ardenne']
        assert ws.cell(3, 1).value == ec.TEXTE_AUCUN_TOURNOI % ec.DATE_RECENSEMENT
        assert ws.cell(4, 1).value == '0 tournoi(s) déclaré(s).'

    def test_contact_et_email_ecrits_au_classeur(self, regenere):
        """Le classeur les porte (col. J et K) ; l'extension ne les lit pas."""
        ws = openpyxl.load_workbook(regenere[0])['Tournois Alsace']
        assert ws.cell(3, 10).value == 'MAITRE JEAN-LOUIS LE MEUR'
        assert ws.cell(3, 11).value == 'jllm.meur@gmail.com'


# ─────────────────────────────────────────────────────────────────────────────
# Aller-retour
# ─────────────────────────────────────────────────────────────────────────────

class TestAllerRetour:

    def test_relecture_rend_le_meme_modele_sauf_les_annotations(self, regenere):
        """Le classeur ne porte plus les annotations : elles n'ont plus de colonne.

        C'est la contrepartie de la décision 3 du §5 — `publication.json` est
        désormais le seul à les détenir. Tout le reste du modèle survit à
        l'aller-retour.
        """
        chemin, avant, _ = regenere
        apres, rapport = ic.lire_classeur(chemin, appliquer_corrections=False)
        assert len(apres) == len(avant) == 45
        champs = [c for c in pub.CHAMPS_DEFAUT if c != 'notes'] + ['tri_key',
                                                                   'cle_source']
        for a, b in zip(avant, apres):
            assert [a[c] for c in champs] == [b[c] for c in champs], a['cle_source']
        assert rapport['notes'] == 0
        assert sum(1 for l in avant if l['notes']) == 12

    def test_relecture_sans_ecart_ni_repechage(self, regenere):
        """Le classeur produit est cohérent : plus de lieu à repêcher, plus d'écart."""
        _, rapport = ic.lire_classeur(regenere[0], appliquer_corrections=False)
        assert rapport['ecarts'] == []
        assert rapport['repechages'] == []
        assert rapport['doublons'] == []

    def test_generation_idempotente(self, regenere, tmp_path):
        """Régénérer depuis la relecture rend le même classeur, cellule pour cellule."""
        lignes, _ = ic.lire_classeur(regenere[0], appliquer_corrections=False)
        for l in lignes:
            l.pop('_source', None)
        deuxieme = str(tmp_path / 'deuxieme.xlsx')
        ec.ecrire_classeur(lignes, deuxieme)
        assert comparer(regenere[0], deuxieme) == []


# ─────────────────────────────────────────────────────────────────────────────
# Outils et garde-fous
# ─────────────────────────────────────────────────────────────────────────────

class TestOutils:

    def test_saison_bascule_au_premier_juillet(self):
        assert ec.saison([{'date_texte': '26-27 sept. 2026'}]) == '2026-2027'
        assert ec.saison([{'date_texte': '3 mai 2027'}]) == '2026-2027'
        assert ec.saison([{'date_texte': '3 juil. 2027'}]) == '2027-2028'

    def test_saison_vide_si_aucune_date_lisible(self):
        """Une saison ne s'invente pas : titre incomplet plutôt que faux."""
        assert ec.saison([{'date_texte': 'un jour'}]) == ''
        assert ec.saison([]) == ''

    def test_bloc_publie_retrouve_la_colonne_c(self):
        for bloc, attendu in [('Parasport', 'Para'),
                              ('Coupe de Lorraine', 'Transverse'),
                              ('Autres', 'Transverse'),
                              ('Épée', 'Épée')]:
            rendu = ec.rendu_chronologique({'bloc': bloc, 'est_tournoi': False})
            assert rendu[2] == attendu
        rendu = ec.rendu_chronologique(
            {'bloc': 'Alsace', 'est_tournoi': True, 'date_texte': '', 'jours': '',
             'lieu': '', 'club': '', 'statut': '', 'armes': []})
        assert rendu[2] == 'Tournoi Alsace'

    def test_ligne_officielle_avant_le_tournoi_de_meme_date(self):
        lignes = [{'date_texte': '1-2 mai 2027', 'est_tournoi': True},
                  {'date_texte': '1-2 mai 2027', 'est_tournoi': False},
                  {'date_texte': '3 oct. 2026', 'est_tournoi': True}]
        assert [(l['date_texte'], l['est_tournoi'])
                for l in ec._ordonnees(lignes)] == [
            ('3 oct. 2026', True), ('1-2 mai 2027', False), ('1-2 mai 2027', True)]

    def test_bloc_inconnu_signale_jamais_deviné(self, tmp_path):
        ligne = pub.ligne_vide(date_texte='3 oct. 2026', bloc='Trampoline',
                               type='ER1')
        rapport = ec.ecrire_classeur([ligne], str(tmp_path / 'c.xlsx'))
        assert len(rapport['anomalies']) == 1
        assert "bloc 'Trampoline' inconnu" in rapport['anomalies'][0]

    def test_bloc_autres_nest_pas_une_anomalie(self, regenere):
        assert regenere[2]['anomalies'] == []

    def test_classeur_ecrase_est_sauvegarde(self, tmp_path):
        chemin = str(tmp_path / 'c.xlsx')
        ec.ecrire_classeur([], chemin)
        rapport = ec.ecrire_classeur([], chemin)
        assert rapport['backup'] and os.path.exists(rapport['backup'])

    def test_generer_lit_publication_json(self, tmp_path, monkeypatch):
        lignes, _ = ic.lire_classeur(CLASSEUR, appliquer_corrections=False)
        for l in lignes:
            l.pop('_source', None)
        monkeypatch.setattr(pub, 'PUBLICATION_FILE', str(tmp_path / 'p.json'))
        pub.save_lignes(lignes)
        rapport = ec.generer(str(tmp_path / 'c.xlsx'))
        assert rapport['lignes'] == 45
        assert os.path.exists(str(tmp_path / 'c.xlsx'))
