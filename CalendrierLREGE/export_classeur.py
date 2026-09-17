# -*- coding: utf-8 -*-
"""Lot 3 — générateur du classeur depuis `data/publication.json`.

Écrit les **12 onglets** que lit l'extension WordPress `calendrier-lrege`
(`Lisez-moi` retiré, décision 4 du §5), à positions de colonnes figées et
dates en texte. C'est l'opération inverse de `import_classeur.py`.

Deux colonnes disparaissent, toutes deux calculées et jamais publiées
(décision 3 du §5) : la colonne I du `Chronologique` — dont seules les
annotations sont conservées, dans le champ `notes` — et la colonne M des
onglets Tournois, qui ne contient que des chevauchements. Le Chronologique
est donc écrit **A→H** et les onglets Tournois **A→L**.

Ce que le générateur dérive, plutôt que de le stocker :

- les **lignes de tournoi du Chronologique**, restitution des onglets
  `Tournois <territoire>` (décision 6 du §5 bis) —
  `import_classeur.rendu_chronologique_tournoi()` ;
- le contenu des **onglets par arme**, depuis le champ `armes[]` des lignes
  officielles (décision 2 du §5) : une ligne transverse (Zone, Master)
  figure sur l'onglet de chacune de ses armes ;
- l'onglet `Transverse`, qui contient exactement les lignes de bloc
  `Coupe de Lorraine` — les `Autres` (AG, CID, Zone, Master) n'y figurent
  pas (vérifié au Lot 2 sur le classeur réel) ;
- les lignes de total des onglets Tournois et le libellé « aucun tournoi ».

`categories` est écrit tel qu'il est saisi, jamais `categories_publiees` :
la forme publiée est un champ calculé, au même titre que `tri_key`
(décision 5 du §5 bis). C'est ce qui permet la régénération à l'identique.

Écarts attendus avec le classeur 2026-2027 d'origine : ils sont énumérés un
par un dans `tests/test_export_classeur.py` (`ECARTS_ATTENDUS`), qui fait
recette. Aucun n'est une perte de saisie — six cellules vides sont
renseignées depuis une autre feuille, deux lignes appliquent une décision déjà
actée, deux cellules tranchent un écart interne au classeur.

Ce que le classeur ne porte plus, en revanche : les **annotations** de la
colonne I du Chronologique. `publication.json` en est désormais le seul
détenteur — c'est la contrepartie de la décision 3 et la raison d'être du
contrôle de recette du Lot 4.

Hors périmètre v1 (§6 du plan) : la mise en forme. Le classeur n'est plus un
document diffusé, c'est un format d'échange vers le site — seules comptent la
structure des onglets et la position des colonnes.
"""
import argparse
import os
import shutil
import sys
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import publication as pub
from import_classeur import (ARMES, LIGNE_ENTETE, ONGLET_CHRONO, ONGLETS_ARME,
                             ONGLETS_TOURNOI, date_en_texte,
                             rendu_chronologique_tournoi)
from texte_lrege import TRI_KEY_ILLISIBLE, tri_key, txt

# ── Structure du classeur produit ────────────────────────────────────────────

#: Date du recensement citée par le libellé « aucun tournoi déclaré ».
DATE_RECENSEMENT = '07/09/2026'

TITRE_CHRONO = ('CALENDRIER CONSOLIDÉ — compétitions officielles '
                '+ tournois de clubs — %s')
TITRE_ARME = '%s — saison %s'
TITRE_TOURNOI = 'TOURNOIS DE CLUBS — %s — saison %s'

#: Chronologique A→H : la colonne I (annotations + chevauchements) est retirée.
ENTETES_CHRONO = ['Date', 'Jour(s)', 'Arme / Bloc', 'Type',
                  'Catégories / Épreuves', 'Lieu', 'Club organisateur', 'Statut']
ENTETES_ARME = ['Date', 'Jour(s)', 'Type', 'Catégories', 'Lieu']
#: Onglets Tournois A→L : la colonne M (chevauchements) est retirée.
ENTETES_TOURNOI = ['Date', 'Jour(s)', 'Club', 'Nom du tournoi', 'Lieu',
                   'Fleuret', 'Épée', 'Sabre', 'Autres épreuves',
                   'Contact', 'E-mail', 'Statut']

TEXTE_AUCUN_TOURNOI = ('Aucun tournoi déclaré pour ce territoire dans le '
                       'recensement au %s.')
TEXTE_TOTAL = '%d tournoi(s) déclaré(s).'

#: Bloc publié -> colonne C du Chronologique. L'inverse de `bloc_publie()`,
#: aux tournois près, dont la ligne est dérivée par
#: `rendu_chronologique_tournoi()` (« Tournoi <territoire> »).
BLOC_CLASSEUR = {'Parasport': 'Para',
                 'Coupe de Lorraine': 'Transverse',
                 'Autres': 'Transverse'}

#: Bloc publié -> onglet qui porte la ligne. `Autres` n'y figure pas : ses
#: lignes (AG, CID, Zone, Master) n'ont pas d'onglet de bloc ; Zone et Master
#: sont portées par les onglets de leurs armes.
ONGLET_DU_BLOC = {'Épée': 'Épée', 'Fleuret': 'Fleuret', 'Sabre': 'Sabre',
                  'Sabre laser': 'Sabre laser', 'Parasport': 'Para',
                  'Coupe de Lorraine': 'Transverse',
                  'Escrime artistique': 'Escrime artistique',
                  'Stage': 'Stages'}

# ── Saison ───────────────────────────────────────────────────────────────────

def saison(lignes):
    """« 2026-2027 », déduite de la première date lisible du lot.

    Bascule au 1er juillet : une date de septembre ouvre la saison, une date
    de mai la ferme. Rend une chaîne vide si aucune date n'est interprétable —
    le titre est alors incomplet, jamais inventé.
    """
    cles = [k for k in (tri_key(l.get('date_texte', '')) for l in lignes)
            if k != TRI_KEY_ILLISIBLE]
    if not cles:
        return ''
    k = min(cles)
    annee, mois = k // 10000, (k // 100) % 100
    debut = annee if mois >= 7 else annee - 1
    return '%d-%d' % (debut, debut + 1)


# ── Rendu d'une ligne ────────────────────────────────────────────────────────

def rendu_chronologique(ligne):
    """Colonnes A→H d'une ligne du Chronologique.

    Une ligne de tournoi n'est pas écrite depuis ses champs : c'est une
    restitution de l'onglet `Tournois <territoire>` (décision 6 du §5 bis).
    """
    if ligne.get('est_tournoi'):
        return rendu_chronologique_tournoi(ligne)
    bloc = txt(ligne.get('bloc', ''))
    return [ligne.get('date_texte', ''), ligne.get('jours', ''),
            BLOC_CLASSEUR.get(bloc, bloc), ligne.get('type', ''),
            ligne.get('categories', ''), ligne.get('lieu', ''),
            ligne.get('club', ''), ligne.get('statut', '')]


def rendu_arme(ligne):
    """Colonnes A→E d'un onglet par arme, de rubrique ou `Transverse`."""
    return [ligne.get('date_texte', ''), ligne.get('jours', ''),
            ligne.get('type', ''), ligne.get('categories', ''),
            ligne.get('lieu', '')]


def rendu_tournoi(ligne):
    """Colonnes A→L d'un onglet Tournois.

    `type` porte le nom du tournoi (col. D) et les catégories sont ventilées
    par arme depuis `armes[]`, comme les lit `import_classeur`.
    """
    cats = {}
    for a in ligne.get('armes') or []:
        if isinstance(a, dict) and txt(a.get('arme', '')):
            cats[txt(a['arme'])] = a.get('categories', '')
    return [ligne.get('date_texte', ''), ligne.get('jours', ''),
            ligne.get('club', ''), ligne.get('type', ''), ligne.get('lieu', ''),
            cats.get('Fleuret', ''), cats.get('Épée', ''), cats.get('Sabre', ''),
            cats.get('Autres', ''),
            ligne.get('contact', ''), ligne.get('email', ''),
            ligne.get('statut', '')]


# ── Sélection et tri ─────────────────────────────────────────────────────────

def _ordonnees(lignes):
    """Lot trié par date, une ligne officielle avant un tournoi de même date.

    C'est l'ordre du Chronologique du classeur, et celui que produit déjà
    `import_classeur.lire_classeur()`. Le tri est stable : à date et nature
    égales, l'ordre d'entrée est conservé.
    """
    return sorted(lignes, key=lambda l: (tri_key(l.get('date_texte', '')),
                                         1 if l.get('est_tournoi') else 0))


def lignes_onglet_arme(lignes, onglet):
    """(lignes du bloc, lignes transverses) d'un onglet par arme.

    Les secondes sont les lignes officielles d'un autre bloc dont `armes[]`
    cite cet onglet — Zone et Master (décision 1 du §5). Le classeur les
    range après les lignes du bloc, chacune séparée par une ligne vide.
    """
    du_bloc, transverses = [], []
    for l in _ordonnees(lignes):
        if l.get('est_tournoi'):
            continue
        if ONGLET_DU_BLOC.get(txt(l.get('bloc', ''))) == onglet:
            du_bloc.append(l)
        elif onglet in ARMES and onglet in (l.get('armes') or []):
            transverses.append(l)
    return du_bloc, transverses


def lignes_onglet_tournoi(lignes, territoire):
    return [l for l in _ordonnees(lignes)
            if l.get('est_tournoi') and txt(l.get('bloc', '')) == territoire]


# ── Écriture ─────────────────────────────────────────────────────────────────

def _ecrire_cellule(ws, r, c, valeur):
    """Écrit une valeur nettoyée. Une valeur vide laisse la cellule vide.

    Les dates sont écrites en **texte** : une cellule datée se relit
    différemment d'un poste à l'autre, et l'extension lit du texte.
    """
    t = txt(date_en_texte(valeur))
    ws.cell(r, c).value = t if t != '' else None


def _ecrire_feuille(ws, titre, entetes, blocs):
    """Titre en ligne 1, en-têtes en ligne 2, puis les blocs séparés d'une ligne vide."""
    _ecrire_cellule(ws, 1, 1, titre)
    for i, e in enumerate(entetes, 1):
        _ecrire_cellule(ws, 2, i, e)
    r = LIGNE_ENTETE + 1
    for i, bloc in enumerate(blocs):
        if i:
            r += 1                       # ligne vide de séparation
        for valeurs in bloc:
            for c, v in enumerate(valeurs, 1):
                _ecrire_cellule(ws, r, c, v)
            r += 1
    return r


def _backup_classeur(chemin):
    """Copie horodatée d'un classeur sur le point d'être écrasé.

    Même principe que `publication.save_lignes()` : rien n'est remplacé sans
    qu'une copie datée reste à côté. Rend le chemin de la copie, ou None s'il
    n'y avait pas de fichier à sauver.
    """
    if not os.path.exists(chemin):
        return None
    dossier = os.path.join(os.path.dirname(os.path.abspath(chemin)), 'backups')
    os.makedirs(dossier, exist_ok=True)
    base, ext = os.path.splitext(os.path.basename(chemin))
    copie = os.path.join(dossier, '%s_%s%s'
                         % (base, datetime.now().strftime('%Y%m%d_%H%M%S_%f'), ext))
    shutil.copy2(chemin, copie)
    return copie


def ecrire_classeur(lignes, chemin, saison_=None,
                    date_recensement=DATE_RECENSEMENT, backup=True):
    """Écrit les 12 onglets du classeur. Rend le rapport d'écriture."""
    lignes = _ordonnees(lignes)
    s = saison(lignes) if saison_ is None else saison_
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    rapport = {'fichier': os.path.basename(chemin), 'saison': s,
               'onglets': [], 'lignes': len(lignes), 'anomalies': []}

    # Chronologique — obligatoire, sinon l'extension refuse l'import.
    ws = wb.create_sheet(ONGLET_CHRONO)
    _ecrire_feuille(ws, TITRE_CHRONO % s, ENTETES_CHRONO,
                    [[rendu_chronologique(l) for l in lignes]])
    rapport['onglets'].append((ONGLET_CHRONO, len(lignes)))

    # Onglets par arme, de rubrique et Transverse.
    for onglet in ONGLETS_ARME:
        ws = wb.create_sheet(onglet)
        du_bloc, transverses = lignes_onglet_arme(lignes, onglet)
        blocs = [[rendu_arme(l) for l in du_bloc]]
        blocs += [[rendu_arme(l)] for l in transverses]
        _ecrire_feuille(ws, TITRE_ARME % (onglet.upper(), s), ENTETES_ARME, blocs)
        rapport['onglets'].append((onglet, len(du_bloc) + len(transverses)))

    # Onglets Tournois, avec leur ligne de total.
    for onglet, territoire in ONGLETS_TOURNOI:
        ws = wb.create_sheet(onglet)
        tournois = lignes_onglet_tournoi(lignes, territoire)
        if tournois:
            blocs = [[rendu_tournoi(l) for l in tournois],
                     [[TEXTE_TOTAL % len(tournois)]]]
        else:
            # Pas de ligne vide de séparation quand le territoire est vide.
            blocs = [[[TEXTE_AUCUN_TOURNOI % date_recensement],
                      [TEXTE_TOTAL % 0]]]
        _ecrire_feuille(ws, TITRE_TOURNOI % (territoire.upper(), s),
                        ENTETES_TOURNOI, blocs)
        rapport['onglets'].append((onglet, len(tournois)))

    # `Autres` est un bloc publié sans onglet propre : AG, CID, Zone et Master
    # ne paraissent qu'au Chronologique — Zone et Master aussi sur les onglets
    # de leurs armes. Ce n'est pas une anomalie.
    blocs_connus = (set(ONGLET_DU_BLOC) | {'Autres'}
                    | {t for _, t in ONGLETS_TOURNOI})
    for l in lignes:
        bloc = txt(l.get('bloc', ''))
        if bloc not in blocs_connus:
            rapport['anomalies'].append(
                '%s %s : bloc %r inconnu — la ligne n\'est écrite qu\'au '
                'Chronologique' % (l.get('date_texte', ''), l.get('type', ''), bloc))

    os.makedirs(os.path.dirname(os.path.abspath(chemin)) or '.', exist_ok=True)
    rapport['backup'] = _backup_classeur(chemin) if backup else None
    tmp = chemin + '.tmp'
    wb.save(tmp)
    os.replace(tmp, chemin)                       # écriture atomique
    return rapport


def generer(chemin, lignes=None, **kw):
    """Écrit le classeur depuis `publication.json`, ou depuis `lignes` fournies."""
    if lignes is None:
        lignes = pub.load_lignes()
    return ecrire_classeur(lignes, chemin, **kw)


def _resume(rapport):
    out = ['%s : saison %s, %d lignes'
           % (rapport['fichier'], rapport['saison'] or '?', rapport['lignes'])]
    out += ['  %-28s %3d ligne(s)' % (nom, n) for nom, n in rapport['onglets']]
    if rapport.get('backup'):
        out.append('\nClasseur précédent sauvegardé : %s'
                   % os.path.basename(rapport['backup']))
    if rapport['anomalies']:
        out.append('\nAnomalies (%d) :' % len(rapport['anomalies']))
        out += ['  - %s' % a for a in rapport['anomalies']]
    return '\n'.join(out)


def main(argv=None):
    p = argparse.ArgumentParser(
        description='Génère le classeur .xlsx depuis publication.json')
    p.add_argument('sortie')
    p.add_argument('--saison', default=None,
                   help='« 2026-2027 » ; déduite des dates si absente')
    p.add_argument('--date-recensement', default=DATE_RECENSEMENT)
    p.add_argument('--sans-backup', action='store_true',
                   help='n\'archive pas le classeur écrasé')
    a = p.parse_args(argv)
    rapport = generer(a.sortie, saison_=a.saison,
                      date_recensement=a.date_recensement,
                      backup=not a.sans_backup)
    print(_resume(rapport))
    return 0


if __name__ == '__main__':
    sys.exit(main())
